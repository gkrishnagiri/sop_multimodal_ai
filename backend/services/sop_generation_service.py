from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openai import OpenAI

from backend.config.settings import settings
from backend.services.job_service import JOBS, update_job


class SopGenerationService:
    """
    MVP 11: Detailed SOP + AI Agent Execution SOP generation from activities JSON.

    Input priority:
        1. data/refined_activities/{job_id}.json
        2. data/activities/{job_id}.json

    Output:
        Detailed SOP for human review:
            data/outputs/{job_id}_sop.json
            data/outputs/{job_id}_sop.md
            data/outputs/{job_id}_sop.docx

        AI Agent SOP for RAG/browser automation:
            data/outputs/{job_id}_agent_sop.json
            data/outputs/{job_id}_agent_sop.md
            data/outputs/{job_id}_agent_sop.docx

    Design principles:
    - Generic across applications, domains, and scenarios
    - No hard-coded business process rules
    - Uses activities/steps/evidence as the source of truth
    - Produces browser-action-oriented instructions for human users and AI agents
    - Does not claim completion unless evidence clearly shows completion
    """

    def __init__(self) -> None:
        self.data_dir = Path(settings.data_dir)
        self.activities_dir = self.data_dir / "activities"
        self.refined_activities_dir = self.data_dir / "refined_activities"
        self.outputs_dir = self.data_dir / "outputs"
        self.outputs_dir.mkdir(parents=True, exist_ok=True)

    def generate_sop_for_job(self, job_id: str) -> dict[str, Any]:
        job = JOBS.get(job_id)

        if not job:
            raise ValueError("Job not found")

        activities_path, activity_source = self._resolve_activity_input_path(job_id)
        activities_data = self._load_json(activities_path)
        compact_input = self._build_compact_llm_input(activities_data)
        observation_summary = compact_input.get("observation_summary", {})

        sop_json = self._generate_sop_with_llm(
            job_id=job_id,
            activities_data=compact_input,
            activity_source=activity_source,
            observation_summary=observation_summary,
        )
        sop_json = self._postprocess_detailed_sop(sop_json, observation_summary)

        agent_sop_json = self._generate_agent_sop_with_llm(
            job_id=job_id,
            activities_data=compact_input,
            detailed_sop=sop_json,
            activity_source=activity_source,
            observation_summary=observation_summary,
        )
        agent_sop_json = self._postprocess_agent_sop(agent_sop_json, observation_summary)

        sop_json_path = self.outputs_dir / f"{job_id}_sop.json"
        sop_md_path = self.outputs_dir / f"{job_id}_sop.md"
        sop_docx_path = self.outputs_dir / f"{job_id}_sop.docx"

        agent_sop_json_path = self.outputs_dir / f"{job_id}_agent_sop.json"
        agent_sop_md_path = self.outputs_dir / f"{job_id}_agent_sop.md"
        agent_sop_docx_path = self.outputs_dir / f"{job_id}_agent_sop.docx"

        self._write_json(sop_json_path, sop_json)
        self._write_json(agent_sop_json_path, agent_sop_json)

        markdown = self._sop_json_to_markdown(sop_json)
        sop_md_path.write_text(markdown, encoding="utf-8")

        agent_markdown = self._agent_sop_json_to_markdown(agent_sop_json)
        agent_sop_md_path.write_text(agent_markdown, encoding="utf-8")

        self._sop_json_to_docx(sop_json, sop_docx_path)
        self._agent_sop_json_to_docx(agent_sop_json, agent_sop_docx_path)

        update_job(
            job_id,
            {
                "sop_json_path": str(sop_json_path),
                "sop_markdown_path": str(sop_md_path),
                "sop_docx_path": str(sop_docx_path),
                "agent_sop_json_path": str(agent_sop_json_path),
                "agent_sop_markdown_path": str(agent_sop_md_path),
                "agent_sop_docx_path": str(agent_sop_docx_path),
                "sop_activity_source": activity_source,
                "status": "sop_generated",
            },
        )

        return {
            "job_id": job_id,
            "status": "sop_generated",
            "activity_source": activity_source,
            "activity_input_path": str(activities_path),
            "sop_json_path": str(sop_json_path),
            "sop_markdown_path": str(sop_md_path),
            "sop_docx_path": str(sop_docx_path),
            "agent_sop_json_path": str(agent_sop_json_path),
            "agent_sop_markdown_path": str(agent_sop_md_path),
            "agent_sop_docx_path": str(agent_sop_docx_path),
            "sop": sop_json,
            "agent_sop": agent_sop_json,
        }

    def get_sop_for_job(self, job_id: str) -> dict[str, Any]:
        sop_json_path = self.outputs_dir / f"{job_id}_sop.json"
        sop_md_path = self.outputs_dir / f"{job_id}_sop.md"
        sop_docx_path = self.outputs_dir / f"{job_id}_sop.docx"

        agent_sop_json_path = self.outputs_dir / f"{job_id}_agent_sop.json"
        agent_sop_md_path = self.outputs_dir / f"{job_id}_agent_sop.md"
        agent_sop_docx_path = self.outputs_dir / f"{job_id}_agent_sop.docx"

        if not sop_json_path.exists():
            raise FileNotFoundError(
                f"SOP file not found: {sop_json_path}. Run generate-sop first."
            )

        sop_json = self._load_json(sop_json_path)
        agent_sop_json = (
            self._load_json(agent_sop_json_path)
            if agent_sop_json_path.exists()
            else {}
        )

        return {
            "job_id": job_id,
            "sop_json_path": str(sop_json_path),
            "sop_markdown_path": str(sop_md_path),
            "sop_docx_path": str(sop_docx_path),
            "agent_sop_json_path": str(agent_sop_json_path),
            "agent_sop_markdown_path": str(agent_sop_md_path),
            "agent_sop_docx_path": str(agent_sop_docx_path),
            "sop": sop_json,
            "agent_sop": agent_sop_json,
            "markdown_available": sop_md_path.exists(),
            "docx_available": sop_docx_path.exists(),
            "agent_markdown_available": agent_sop_md_path.exists(),
            "agent_docx_available": agent_sop_docx_path.exists(),
        }

    def _resolve_activity_input_path(self, job_id: str) -> tuple[Path, str]:
        refined_path = self.refined_activities_dir / f"{job_id}.json"
        if refined_path.exists():
            return refined_path, "refined_activities_json"

        activities_path = self.activities_dir / f"{job_id}.json"
        if activities_path.exists():
            return activities_path, "activities_json"

        raise FileNotFoundError(
            f"No activity input found for job_id={job_id}. "
            "Run detect-activities first. For better SOP quality, run refine-activities after that."
        )

    def _generate_sop_with_llm(
        self,
        job_id: str,
        activities_data: dict[str, Any],
        activity_source: str,
        observation_summary: dict[str, Any],
    ) -> dict[str, Any]:
        if not settings.openai_api_key:
            raise ValueError("OPENAI_API_KEY is missing. Add it to your .env file.")

        client = OpenAI(api_key=settings.openai_api_key)

        response = client.chat.completions.create(
            model=settings.llm_model,
            temperature=0.15,
            response_format={
                "type": "json_schema",
                "json_schema": {
                    "name": "agent_ready_sop_generation_result",
                    "strict": True,
                    "schema": self._sop_json_schema(),
                },
            },
            messages=[
                {
                    "role": "system",
                    "content": self._system_prompt(),
                },
                {
                    "role": "user",
                    "content": self._user_prompt(
                        job_id=job_id,
                        activities_data=activities_data,
                        activity_source=activity_source,
                        observation_summary=observation_summary,
                    ),
                },
            ],
        )

        content = response.choices[0].message.content

        if not content:
            raise RuntimeError("LLM returned empty SOP response.")

        try:
            sop = json.loads(content)
        except json.JSONDecodeError as exc:
            raise RuntimeError(f"LLM returned invalid JSON: {exc}") from exc

        sop["job_id"] = job_id
        sop["metadata"] = {
            "llm_model": settings.llm_model,
            "generator": "generic_agent_ready_sop_generator_v2",
            "source": activity_source,
        }

        return sop

    def _generate_agent_sop_with_llm(
        self,
        job_id: str,
        activities_data: dict[str, Any],
        detailed_sop: dict[str, Any],
        activity_source: str,
        observation_summary: dict[str, Any],
    ) -> dict[str, Any]:
        if not settings.openai_api_key:
            raise ValueError("OPENAI_API_KEY is missing. Add it to your .env file.")

        client = OpenAI(api_key=settings.openai_api_key)

        detailed_sop_summary = self._build_detailed_sop_summary(detailed_sop)

        response = client.chat.completions.create(
            model=settings.llm_model,
            temperature=0.12,
            response_format={
                "type": "json_schema",
                "json_schema": {
                    "name": "generic_agent_execution_sop_result",
                    "strict": True,
                    "schema": self._agent_sop_json_schema(),
                },
            },
            messages=[
                {
                    "role": "system",
                    "content": self._agent_system_prompt(),
                },
                {
                    "role": "user",
                    "content": self._agent_user_prompt(
                        job_id=job_id,
                        activities_data=activities_data,
                        detailed_sop_summary=detailed_sop_summary,
                        activity_source=activity_source,
                        observation_summary=observation_summary,
                    ),
                },
            ],
        )

        content = response.choices[0].message.content

        if not content:
            raise RuntimeError("LLM returned empty Agent SOP response.")

        try:
            agent_sop = json.loads(content)
        except json.JSONDecodeError as exc:
            raise RuntimeError(f"LLM returned invalid Agent SOP JSON: {exc}") from exc

        agent_sop["job_id"] = job_id
        agent_sop["metadata"] = {
            "llm_model": settings.llm_model,
            "generator": "generic_agent_execution_sop_generator_v2",
            "source": activity_source,
            "derived_from": "detailed_sop_and_activity_evidence",
        }

        return agent_sop

    def _system_prompt(self) -> str:
        return """
You are an expert SOP analyst and browser automation workflow designer.

Your task is to convert multimodal workflow activity data into an Agent-ready Standard Operating Procedure.

The SOP will be used by:
1. A human beginner who needs clear browser instructions.
2. An AI Agent that will use the SOP through RAG, create an execution plan, and perform browser actions.

Important rules:
1. Do not assume the application, business domain, user intent, final outcome, or process completion beyond the evidence provided.
2. Do not invent steps, field values, buttons, screens, outcomes, or final completion states that are not supported by activities, speech, screen text, timestamps, or frame references.
3. Do not hard-code any website, application, company, or scenario-specific behavior.
4. Keep the SOP generic enough for any similarly recorded browser/application workflow, while using observed terminology when clearly supported by evidence.
5. Every executable step must be written as a clear browser/UI instruction.
6. Prefer imperative action verbs: Open, Click, Type, Select, Check, Review, Wait, Scroll, Confirm, Enter, Choose, Search, Proceed.
7. Avoid vague instructions such as "Use filters", "Navigate to relevant page", "Configure visible settings", or "Select appropriate option" unless the specific target cannot be inferred.
8. If a field value is not visible or not safely reusable, write "Enter the required value for <field name>".
9. Avoid raw personal data. Use generic terms such as "saved passenger", "registered contact details", "selected record", "user account", "selected item", or "configured option".
10. Distinguish observed action from reusable instruction. The reusable instruction should be executable, but it must remain evidence-grounded.
11. Preserve traceability by including evidence timestamps and frame paths in each step.
12. Do not include finalization words such as "finalize", "complete", "submit", "approve", "pay", "save", "close", or "confirm" as completed actions unless evidence explicitly shows that completed state.
13. If recording stops before final payment/submission/approval/completion, the observed_end_state and completion criteria must say that the workflow reached the last observed review/pre-submit/payment-entry screen.
14. Expected results must describe immediate visible screen results, not downstream business outcomes.
15. Validation must describe what the agent or human should check on screen before continuing.
16. Do not create one step per video timestamp, frame, or scroll event. Summarize repeated actions into one reusable step with a loop or repeat note.
17. Return only valid JSON matching the provided schema.
""".strip()

    def _user_prompt(
        self,
        job_id: str,
        activities_data: dict[str, Any],
        activity_source: str,
        observation_summary: dict[str, Any],
    ) -> str:
        return f"""
Generate an Agent-ready browser execution SOP from the following generic activity data.

Job ID:
{job_id}

Activity source:
{activity_source}

Activity data:
{json.dumps(activities_data, ensure_ascii=False, indent=2)}

Observation summary extracted from OCR/screen text/speech:
{json.dumps(observation_summary, ensure_ascii=False, indent=2)}

Output requirements:
- Produce a beginner-friendly SOP guide.
- Produce browser-action-oriented steps suitable for a RAG-based AI Agent.
- Each step must have:
  - instruction
  - ui_action
  - action_type
  - target_ui_element
  - input_value_description
  - expected_result
  - validation
  - condition
  - fallback_or_note
  - evidence
- Use explicit actions when evidence supports them.
- Use observed_urls, observed_entry_point, observed_pages, access_terms, observed_buttons, observed_input_fields, observed_menu_items, observed_form_sections, and observed_result_sections from the observation summary when present. Do not say URLs or entry points are not observed if observation_summary contains them.
- Merge repeated adjacent review/check/filter/scroll actions into a single reusable instruction when later evidence adds no new UI action information.
- If many frames show the same page/state, do not generate one SOP step per frame.
- If OCR/screen text is weak, use speech and activity evidence conservatively.
- If the exact UI element is unknown, use a generic but useful target such as "the visible Search button", "the journey date field", "the selected option", or "the payment method section".
- Keep evidence timestamps and frame paths.
- Do not invent missing field values.
- Do not claim final booking/payment/submission/completion unless evidence explicitly proves it.
- If final state is only payment review or payment details entry, stop there.

Return JSON only.
""".strip()

    def _agent_system_prompt(self) -> str:
        return """
You are an expert browser automation SOP designer.

Your task is to convert observed multimodal workflow data into a concise Agent Execution SOP.

This Agent SOP will be used by a RAG-based AI agent that creates a browser execution plan and performs actions using tools such as Playwright.

The document must be generic and reusable:
- It must work for any application, website, portal, or UI workflow recorded in a video.
- Do not hard-code IRCTC, Vendor Admin, or any other sample-specific rule unless it is directly observed in the input.
- Use observed application terminology only when supported by evidence.
- Convert observed behavior into task-oriented reusable procedures.

Writing rules:
1. Organize the SOP like an application guide, not like an audit report.
2. Explain the observed application or workflow purpose at a high level.
3. Include access/navigation details only if observed. OCR URLs, browser address text, page routes, login/register labels, and screen headings count as observed access/navigation details. If observed_urls or observed_entry_point are provided, use them.
4. Include prerequisites only if observed or safely inferable from the workflow.
5. Create a task catalog from broad observed activities.
6. For each task, write short one-action steps. Each task should normally have 3 to 12 steps unless the evidence truly shows more unique user actions.
7. Every step must have exactly one primary browser/UI action.
8. Every step must include a success_check that the agent can verify on screen.
9. Use simple imperative verbs: Open, Click, Enter, Select, Check, Review, Wait, Scroll, Search, Choose, Confirm.
10. Do not include local screenshot paths, timestamps, transcript snippets, or evidence blocks in the Agent SOP.
11. Do not include long validation tables in the Agent SOP.
12. Do not invent credentials, URLs, field values, train names, user names, IDs, passwords, amounts, or final outcomes.
13. If a value is task-specific, write "Enter the required value for <field name>".
14. If a target is not exact but visible context exists, use a practical generic target like "the visible Search button" or "the selected record row".
15. Do not claim payment, submission, approval, deletion, creation, or final completion unless evidence explicitly proves it.
16. Return only valid JSON matching the provided schema.
""".strip()

    def _agent_user_prompt(
        self,
        job_id: str,
        activities_data: dict[str, Any],
        detailed_sop_summary: dict[str, Any],
        activity_source: str,
        observation_summary: dict[str, Any],
    ) -> str:
        return f"""
Generate a concise AI Agent Execution SOP from this workflow evidence.

Job ID:
{job_id}

Activity source:
{activity_source}

Compact activity data:
{json.dumps(activities_data, ensure_ascii=False, indent=2)}

Detailed SOP summary:
{json.dumps(detailed_sop_summary, ensure_ascii=False, indent=2)}

Observation summary extracted from OCR/screen text/speech:
{json.dumps(observation_summary, ensure_ascii=False, indent=2)}

Output requirements:
- Create a generic application/task guide suitable for a RAG-based browser automation agent.
- Include sections for application overview, access/navigation, prerequisites, general rules, task catalog, and limitations.
- Translate broad observed activities into reusable tasks.
- Use observed_urls and observed_entry_point as the application entry/navigation details when present. Do not say "Not observed in the recording" for URL or entry point if observation_summary contains URL/page evidence.
- Optimize and summarize repeated review/check/filter/scroll steps. Do not repeat identical instructions and success checks.
- Prefer loop-style steps such as "For each visible record...", "For each available option...", or "Repeat until all visible items are reviewed" when evidence shows repeated review of similar UI items.
- For every task, provide required inputs when applicable.
- For every task, provide expected_agent_output. This tells the execution agent what final information to return for review/search/recommendation tasks.
- Prefer observed UI labels from observation_summary when available:
  - observed_input_fields
  - observed_buttons
  - observed_menu_items
  - observed_form_sections
  - observed_result_sections
- Do not use vague phrases such as "relevant page", "visible criteria", or "displayed information" when an observed URL, page, field, button, section, or menu label is available.
- For every step:
  - Write one simple browser/UI action.
  - Include a success_check that tells the agent what to verify before continuing.
  - Include action_type and target_ui_element.
  - Keep notes short.
- Do not include evidence paths, timestamps, or screenshots in the Agent SOP.
- Do not generate a table-heavy audit document.
- Do not invent missing values. Use placeholders like "Enter the required value for <field name>" when needed.
- If the recording only shows a review/pre-submit/payment-entry state, reflect that limitation.

Return JSON only.
""".strip()

    def _build_detailed_sop_summary(self, sop: dict[str, Any]) -> dict[str, Any]:
        activities_summary = []

        for activity in sop.get("activities", []):
            step_summaries = []

            for step in activity.get("steps", [])[:20]:
                step_summaries.append(
                    {
                        "step_number": step.get("step_number"),
                        "instruction": self._clean_text(step.get("instruction", "")),
                        "ui_action": self._clean_text(step.get("ui_action", "")),
                        "action_type": step.get("action_type", ""),
                        "target_ui_element": self._clean_text(step.get("target_ui_element", "")),
                        "expected_result": self._clean_text(step.get("expected_result", "")),
                        "validation": self._clean_text(step.get("validation", "")),
                    }
                )

            activities_summary.append(
                {
                    "activity_id": activity.get("activity_id", ""),
                    "name": self._clean_text(activity.get("name", "")),
                    "description": self._clean_text(activity.get("description", "")),
                    "steps": step_summaries,
                }
            )

        return {
            "document_title": self._clean_text(sop.get("document_title", "")),
            "overview": self._clean_text(sop.get("overview", "")),
            "scope": self._clean_text(sop.get("scope", "")),
            "observed_end_state": self._clean_text(sop.get("observed_end_state", "")),
            "business_process_description": self._clean_text(
                sop.get("business_process_description", "")
            ),
            "prerequisites": [
                self._clean_text(item)
                for item in sop.get("prerequisites", [])
            ],
            "general_rules": [
                self._clean_text(item)
                for item in sop.get("general_rules", [])
            ],
            "activities": activities_summary,
            "completion_criteria": [
                self._clean_text(item)
                for item in sop.get("completion_criteria", [])
            ],
        }

    def _agent_sop_json_schema(self) -> dict[str, Any]:
        step_schema = {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "step_number": {"type": "integer"},
                "instruction": {"type": "string"},
                "success_check": {"type": "string"},
                "action_type": {
                    "type": "string",
                    "enum": [
                        "open",
                        "click",
                        "type",
                        "select",
                        "check",
                        "review",
                        "wait",
                        "scroll",
                        "search",
                        "choose",
                        "confirm",
                        "conditional",
                        "other",
                    ],
                },
                "target_ui_element": {"type": "string"},
                "notes": {"type": "string"},
            },
            "required": [
                "step_number",
                "instruction",
                "success_check",
                "action_type",
                "target_ui_element",
                "notes",
            ],
        }

        task_schema = {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "task_name": {"type": "string"},
                "task_purpose": {"type": "string"},
                "required_inputs": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "steps": {
                    "type": "array",
                    "items": step_schema,
                },
                "completion_criteria": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "expected_agent_output": {
                    "type": "array",
                    "items": {"type": "string"},
                },
            },
            "required": [
                "task_name",
                "task_purpose",
                "required_inputs",
                "steps",
                "completion_criteria",
                "expected_agent_output",
            ],
        }

        return {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "job_id": {"type": "string"},
                "document_title": {"type": "string"},
                "classification": {"type": "string"},
                "application_overview": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "application_name": {"type": "string"},
                        "application_type": {"type": "string"},
                        "observed_url_or_entry_point": {"type": "string"},
                        "purpose": {"type": "string"},
                    },
                    "required": [
                        "application_name",
                        "application_type",
                        "observed_url_or_entry_point",
                        "purpose",
                    ],
                },
                "prerequisites": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "access_and_navigation": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "general_rules": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "task_catalog": {
                    "type": "array",
                    "items": task_schema,
                },
                "limitations": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "agent_execution_notes": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "metadata": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "llm_model": {"type": "string"},
                        "generator": {"type": "string"},
                        "source": {"type": "string"},
                        "derived_from": {"type": "string"},
                    },
                    "required": [
                        "llm_model",
                        "generator",
                        "source",
                        "derived_from",
                    ],
                },
            },
            "required": [
                "job_id",
                "document_title",
                "classification",
                "application_overview",
                "prerequisites",
                "access_and_navigation",
                "general_rules",
                "task_catalog",
                "limitations",
                "agent_execution_notes",
                "metadata",
            ],
        }

    def _agent_sop_json_to_markdown(self, agent_sop: dict[str, Any]) -> str:
        lines: list[str] = []

        lines.append(f"# {agent_sop.get('document_title', 'AI Agent Execution SOP')}")
        lines.append("")
        lines.append(f"Classification: {agent_sop.get('classification', 'Internal')}")
        lines.append("")

        overview = agent_sop.get("application_overview", {})
        lines.append("## 1. Application Overview")
        lines.append(f"- Application name: {overview.get('application_name', 'Not observed')}")
        lines.append(f"- Application type: {overview.get('application_type', 'Not observed')}")
        lines.append(f"- Observed URL or entry point: {overview.get('observed_url_or_entry_point', 'Not observed')}")
        lines.append(f"- Purpose: {overview.get('purpose', 'Not specified')}")
        lines.append("")

        self._append_markdown_list(lines, "2. Access and Navigation", agent_sop.get("access_and_navigation", []))
        self._append_markdown_list(lines, "3. Prerequisites", agent_sop.get("prerequisites", []))
        self._append_markdown_list(lines, "4. General Rules Before Execution", agent_sop.get("general_rules", []))

        lines.append("## 5. Task Catalog")
        lines.append("")

        for index, task in enumerate(agent_sop.get("task_catalog", []), start=1):
            lines.append(f"### 5.{index}. {task.get('task_name', 'Task')}")
            lines.append("")
            lines.append(f"Purpose: {task.get('task_purpose', '')}")
            lines.append("")

            required_inputs = task.get("required_inputs", [])
            if required_inputs:
                lines.append("Required inputs:")
                for item in required_inputs:
                    lines.append(f"- {item}")
                lines.append("")

            lines.append("Steps:")
            for step in task.get("steps", []):
                lines.append(f"Step {step.get('step_number')}: {step.get('instruction', '')}")
                lines.append(f"Success check: {step.get('success_check', '')}")

                notes = step.get("notes", "")
                if notes:
                    lines.append(f"Note: {notes}")

                lines.append("")

            completion_criteria = task.get("completion_criteria", [])
            if completion_criteria:
                lines.append("Completion criteria:")
                for item in completion_criteria:
                    lines.append(f"- {item}")
                lines.append("")

            expected_agent_output = task.get("expected_agent_output", [])
            if expected_agent_output:
                lines.append("Expected agent output:")
                for item in expected_agent_output:
                    lines.append(f"- {item}")
                lines.append("")

        self._append_markdown_list(lines, "6. Limitations", agent_sop.get("limitations", []))
        self._append_markdown_list(lines, "7. Agent Execution Notes", agent_sop.get("agent_execution_notes", []))

        return "\n".join(lines).strip() + "\n"

    def _agent_sop_json_to_docx(self, agent_sop: dict[str, Any], output_path: Path) -> None:
        document = Document()
        self._configure_docx(document)

        classification = agent_sop.get("classification") or "Internal"
        document.add_paragraph(f"Classification: {classification}")

        title = agent_sop.get("document_title") or "AI Agent Execution SOP"
        title_paragraph = document.add_paragraph()
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_paragraph.add_run(title)
        title_run.bold = True
        title_run.font.size = Pt(18)

        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_run = subtitle.add_run("AI Agent Execution Guide")
        subtitle_run.italic = True
        subtitle_run.font.size = Pt(11)

        document.add_paragraph("-" * 80)

        overview = agent_sop.get("application_overview", {})
        document.add_heading("1. APPLICATION OVERVIEW", level=1)
        self._add_labeled_paragraph(document, "Application name", overview.get("application_name", "Not observed"))
        self._add_labeled_paragraph(document, "Application type", overview.get("application_type", "Not observed"))
        self._add_labeled_paragraph(document, "Observed URL or entry point", overview.get("observed_url_or_entry_point", "Not observed"))
        self._add_labeled_paragraph(document, "Purpose", overview.get("purpose", "Not specified"))

        self._add_bullet_section(document, 2, "ACCESS AND NAVIGATION", agent_sop.get("access_and_navigation", []))
        self._add_bullet_section(document, 3, "PREREQUISITES", agent_sop.get("prerequisites", []))
        self._add_bullet_section(document, 4, "GENERAL RULES BEFORE EXECUTION", agent_sop.get("general_rules", []))

        document.add_heading("5. TASK CATALOG", level=1)

        for task_index, task in enumerate(agent_sop.get("task_catalog", []), start=1):
            document.add_heading(f"5.{task_index}. {task.get('task_name', 'Task')}", level=2)

            purpose = task.get("task_purpose")
            if purpose:
                self._add_labeled_paragraph(document, "Purpose", purpose)

            required_inputs = task.get("required_inputs", [])
            if required_inputs:
                document.add_paragraph("Required inputs:")
                for item in required_inputs:
                    document.add_paragraph(str(item), style="List Bullet")

            document.add_paragraph("Steps:")

            for step in task.get("steps", []):
                instruction = step.get("instruction", "")
                success_check = step.get("success_check", "")

                paragraph = document.add_paragraph()
                run = paragraph.add_run(f"Step {step.get('step_number')}: {instruction}")
                run.bold = True

                success_paragraph = document.add_paragraph()
                success_run = success_paragraph.add_run("Success check: ")
                success_run.bold = True
                success_paragraph.add_run(str(success_check))

                notes = step.get("notes", "")
                if notes:
                    note_paragraph = document.add_paragraph()
                    note_run = note_paragraph.add_run("Note: ")
                    note_run.bold = True
                    note_paragraph.add_run(str(notes))

            completion_criteria = task.get("completion_criteria", [])
            if completion_criteria:
                document.add_paragraph("Task completion criteria:")
                for item in completion_criteria:
                    document.add_paragraph(str(item), style="List Bullet")

            expected_agent_output = task.get("expected_agent_output", [])
            if expected_agent_output:
                document.add_paragraph("Expected agent output:")
                for item in expected_agent_output:
                    document.add_paragraph(str(item), style="List Bullet")

        self._add_bullet_section(document, 6, "LIMITATIONS", agent_sop.get("limitations", []))
        self._add_bullet_section(document, 7, "AGENT EXECUTION NOTES", agent_sop.get("agent_execution_notes", []))

        document.add_paragraph("-" * 80)
        document.add_paragraph("END OF DOCUMENT")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)

    def _add_labeled_paragraph(self, document: Document, label: str, value: Any) -> None:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"{label}: ")
        run.bold = True
        paragraph.add_run(str(value or "Not specified."))

    def _build_compact_llm_input(self, activities_data: dict[str, Any]) -> dict[str, Any]:
        compact_activities = []

        for activity in activities_data.get("activities", []):
            compact_steps = []

            for step in activity.get("steps", []):
                evidence = step.get("evidence", {})
                if not isinstance(evidence, dict):
                    evidence = {}

                screen_text = (
                    step.get("screen_text_sample")
                    or evidence.get("screen_text")
                    or []
                )

                speech = (
                    step.get("speech")
                    or evidence.get("speech_summary")
                    or ""
                )

                compact_steps.append(
                    {
                        "step_number": step.get("step_number"),
                        "start_seconds": step.get("start_seconds"),
                        "end_seconds": step.get("end_seconds"),
                        "intent": step.get("intent", ""),
                        "instruction": self._clean_text(step.get("instruction", "")),
                        "ui_action": self._clean_text(step.get("ui_action", "")),
                        "expected_result": self._clean_text(step.get("expected_result", "")),
                        "speech": self._clean_text(speech),
                        "screen_text_sample": self._dedupe_keep_order(screen_text)[:12],
                        "frame_path": step.get("frame_path") or evidence.get("frame_path", ""),
                        "source_step_refs": step.get("source_step_refs", []),
                    }
                )

            evidence_block = activity.get("evidence", {})
            if not isinstance(evidence_block, dict):
                evidence_block = {}

            compact_activities.append(
                {
                    "activity_id": activity.get("activity_id"),
                    "name": self._clean_text(activity.get("name", "")),
                    "description": self._clean_text(activity.get("description", "")),
                    "start_seconds": activity.get("start_seconds"),
                    "end_seconds": activity.get("end_seconds"),
                    "duration_seconds": activity.get("duration_seconds"),
                    "dominant_intent": activity.get("dominant_intent", ""),
                    "source_activity_ids": activity.get("source_activity_ids", []),
                    "evidence_summary": self._clean_text(activity.get("evidence_summary", "")),
                    "evidence": {
                        "speech_samples": [
                            self._clean_text(item)
                            for item in evidence_block.get("speech_samples", [])[:5]
                        ],
                        "screen_text_samples": self._dedupe_keep_order(
                            evidence_block.get("screen_text_samples", [])
                        )[:12],
                        "frame_paths": evidence_block.get("frame_paths", [])[:3],
                    },
                    "steps": compact_steps,
                }
            )

        compact_payload = {
            "job_id": activities_data.get("job_id"),
            "activity_count": len(compact_activities),
            "activities": compact_activities,
            "metadata": activities_data.get("metadata", {}),
        }
        compact_payload["observation_summary"] = self._build_observation_summary(compact_payload)
        return compact_payload

    def _build_observation_summary(self, compact_payload: dict[str, Any]) -> dict[str, Any]:
        """Extract deterministic application/navigation/UI signals before LLM generation.

        This is intentionally generic. It does not contain application-specific
        workflow rules. It promotes what was actually observed in OCR/screen
        text/speech into a compact signal set that the SOP prompts can use.
        """

        all_screen_text: list[str] = []
        all_speech: list[str] = []
        page_states: list[str] = []
        url_candidates: list[str] = []

        for activity in compact_payload.get("activities", []):
            evidence = activity.get("evidence", {}) if isinstance(activity.get("evidence"), dict) else {}

            for item in evidence.get("screen_text_samples", []) or []:
                all_screen_text.append(self._clean_text(item))

            for item in evidence.get("speech_samples", []) or []:
                all_speech.append(self._clean_text(item))

            for step in activity.get("steps", []) or []:
                for item in step.get("screen_text_sample", []) or []:
                    all_screen_text.append(self._clean_text(item))

                speech = self._clean_text(step.get("speech", ""))
                if speech:
                    all_speech.append(speech)

        clean_text = self._dedupe_keep_order([item for item in all_screen_text if item])
        clean_speech = self._dedupe_keep_order([item for item in all_speech if item])

        for text in clean_text:
            for url in self._extract_url_candidates(text):
                url_candidates.append(url)

        url_candidates = self._dedupe_keep_order(url_candidates)

        for url in url_candidates:
            page = self._url_to_page_state(url)
            if page:
                page_states.append(page)

        observed_buttons = self._extract_observed_buttons(clean_text)
        observed_input_fields = self._extract_observed_input_fields(clean_text)
        observed_menu_items = self._extract_observed_menu_items(clean_text)
        observed_form_sections = self._extract_observed_form_sections(clean_text)
        observed_result_sections = self._extract_observed_result_sections(clean_text)

        access_terms = []
        access_keywords = [
            "login", "log in", "register", "sign in", "homepage", "home page",
            "welcome", "dashboard", "my account", "menu", "profile", "search",
            "account", "portal", "home"
        ]

        for text in clean_text:
            lower = text.lower()
            if any(keyword in lower for keyword in access_keywords):
                access_terms.append(text)

        entry_point = url_candidates[0] if url_candidates else ""

        return {
            "observed_entry_point": entry_point,
            "observed_urls": url_candidates[:12],
            "observed_pages": self._dedupe_keep_order(page_states)[:12],
            "access_terms": self._dedupe_keep_order(access_terms)[:20],
            "observed_buttons": observed_buttons[:30],
            "observed_input_fields": observed_input_fields[:30],
            "observed_menu_items": observed_menu_items[:30],
            "observed_form_sections": observed_form_sections[:20],
            "observed_result_sections": observed_result_sections[:20],
            "representative_screen_text": clean_text[:30],
            "representative_speech": clean_speech[:12],
        }

    def _extract_observed_buttons(self, texts: list[str]) -> list[str]:
        action_words = [
            "login", "log in", "register", "sign in", "search", "submit", "save",
            "continue", "proceed", "confirm", "cancel", "ok", "yes", "no", "apply",
            "filter", "reset", "next", "back", "add", "delete", "edit", "update",
            "select", "choose", "book", "review", "pay"
        ]

        buttons = []
        for text in texts:
            clean = self._clean_text(text)
            lower = clean.lower()
            if not clean or len(clean) > 60:
                continue
            if any(word == lower or word in lower for word in action_words):
                buttons.append(clean)

        return self._dedupe_keep_order(buttons)

    def _extract_observed_input_fields(self, texts: list[str]) -> list[str]:
        field_keywords = [
            "name", "username", "user name", "password", "email", "mail", "mobile",
            "phone", "age", "gender", "date", "time", "from", "to", "source",
            "destination", "address", "city", "state", "country", "pin", "zip",
            "amount", "number", "id", "code", "class", "type", "category",
            "preference", "option", "details", "information"
        ]

        fields = []
        for text in texts:
            clean = self._clean_text(text)
            lower = clean.lower()
            if not clean or len(clean) > 80:
                continue
            if any(keyword == lower or keyword in lower for keyword in field_keywords):
                fields.append(clean)

        return self._dedupe_keep_order(fields)

    def _extract_observed_menu_items(self, texts: list[str]) -> list[str]:
        menu_keywords = [
            "menu", "account", "profile", "settings", "dashboard", "home",
            "reports", "admin", "users", "roles", "permissions", "search",
            "transactions", "orders", "booking", "history", "help"
        ]

        items = []
        for text in texts:
            clean = self._clean_text(text)
            lower = clean.lower()
            if not clean or len(clean) > 60:
                continue
            if any(keyword == lower or keyword in lower for keyword in menu_keywords):
                items.append(clean)

        return self._dedupe_keep_order(items)

    def _extract_observed_form_sections(self, texts: list[str]) -> list[str]:
        section_keywords = [
            "details", "information", "input", "preferences", "settings",
            "options", "payment", "billing", "shipping", "profile", "address",
            "contact", "passenger", "customer", "user", "role", "permission",
            "configuration", "summary", "review"
        ]

        sections = []
        for text in texts:
            clean = self._clean_text(text)
            lower = clean.lower()
            if not clean or len(clean) > 90:
                continue
            if any(keyword in lower for keyword in section_keywords):
                sections.append(clean)

        return self._dedupe_keep_order(sections)

    def _extract_observed_result_sections(self, texts: list[str]) -> list[str]:
        result_keywords = [
            "result", "results", "list", "table", "records", "status",
            "available", "availability", "summary", "row", "option",
            "options", "search results", "items", "entries"
        ]

        sections = []
        for text in texts:
            clean = self._clean_text(text)
            lower = clean.lower()
            if not clean or len(clean) > 100:
                continue
            if any(keyword in lower for keyword in result_keywords):
                sections.append(clean)

        return self._dedupe_keep_order(sections)

    def _extract_url_candidates(self, text: str) -> list[str]:
        text = self._clean_text(text)
        if not text:
            return []

        matches = re.findall(
            r"(?:https?://)?(?:www\.)?[A-Za-z0-9.-]+\.(?:com|in|org|net|io)(?:/[A-Za-z0-9._~:/?#\[\]@!$&'()*+,;=%-]*)?",
            text,
        )
        return [match.rstrip(".,;:)") for match in matches]

    def _url_to_page_state(self, url: str) -> str:
        clean = re.sub(r"^https?://", "", url.lower()).strip()
        clean = re.sub(r"\?.*$", "", clean)
        clean = re.sub(r"#.*$", "", clean)
        return clean.strip("/")

    def _postprocess_detailed_sop(
        self,
        sop: dict[str, Any],
        observation_summary: dict[str, Any],
    ) -> dict[str, Any]:
        """Deterministic cleanup after LLM generation for detailed SOP quality."""
        sop = dict(sop)
        activities = []

        for activity in sop.get("activities", []) or []:
            if not isinstance(activity, dict):
                continue
            cleaned_activity = dict(activity)
            cleaned_activity["steps"] = self._compress_detailed_steps(
                cleaned_activity.get("steps", []) or []
            )
            activities.append(cleaned_activity)

        sop["activities"] = activities
        sop["completion_criteria"] = self._clean_completion_criteria(
            sop.get("completion_criteria", []) or [],
            sop,
        )

        metadata = sop.get("metadata") if isinstance(sop.get("metadata"), dict) else {}
        metadata["postprocessor"] = "mvp11_2_execution_quality_hardening"
        metadata["observed_entry_point"] = observation_summary.get("observed_entry_point", "")
        sop["metadata"] = metadata

        return sop

    def _postprocess_agent_sop(
        self,
        agent_sop: dict[str, Any],
        observation_summary: dict[str, Any],
    ) -> dict[str, Any]:
        """Deterministic cleanup after LLM generation for Agent SOP quality."""
        agent_sop = dict(agent_sop)
        overview = agent_sop.get("application_overview")
        if not isinstance(overview, dict):
            overview = {}

        observed_entry_point = self._clean_text(observation_summary.get("observed_entry_point", ""))
        observed_urls = observation_summary.get("observed_urls", []) or []
        observed_pages = observation_summary.get("observed_pages", []) or []
        access_terms = observation_summary.get("access_terms", []) or []

        current_entry = self._clean_text(overview.get("observed_url_or_entry_point", ""))
        if observed_entry_point and (
            not current_entry
            or "not observed" in current_entry.lower()
            or current_entry.lower() in {"none", "n/a", "not specified"}
        ):
            overview["observed_url_or_entry_point"] = observed_entry_point

        app_type = self._clean_text(overview.get("application_type", ""))
        if not app_type or "not observed" in app_type.lower():
            overview["application_type"] = "Web application" if observed_urls else "Application workflow"

        agent_sop["application_overview"] = overview

        access_and_navigation = [
            self._clean_text(item)
            for item in agent_sop.get("access_and_navigation", []) or []
            if self._clean_text(item)
        ]
        access_joined = " ".join(access_and_navigation).lower()

        if observed_entry_point and observed_entry_point.lower() not in access_joined:
            access_and_navigation.insert(0, f"Open the observed entry point: {observed_entry_point}")

        for page in observed_pages[:5]:
            if page and page.lower() not in " ".join(access_and_navigation).lower():
                access_and_navigation.append(f"Observed application page/route: {page}")

        for term in access_terms[:5]:
            if term and term.lower() not in " ".join(access_and_navigation).lower():
                access_and_navigation.append(f"Observed access/navigation cue: {term}")

        agent_sop["access_and_navigation"] = self._dedupe_keep_order(access_and_navigation)

        cleaned_tasks = []
        for task in agent_sop.get("task_catalog", []) or []:
            if not isinstance(task, dict):
                continue
            cleaned_task = dict(task)
            cleaned_task["steps"] = self._compress_agent_steps(cleaned_task.get("steps", []) or [])
            cleaned_task["steps"] = self._improve_agent_steps_with_observations(
                cleaned_task.get("steps", []) or [],
                cleaned_task,
                observation_summary,
            )
            cleaned_task["expected_agent_output"] = self._build_expected_agent_output(
                cleaned_task,
                observation_summary,
            )
            cleaned_task["completion_criteria"] = self._clean_completion_criteria(
                cleaned_task.get("completion_criteria", []) or [],
                agent_sop,
            )
            cleaned_tasks.append(cleaned_task)

        agent_sop["task_catalog"] = cleaned_tasks

        metadata = agent_sop.get("metadata") if isinstance(agent_sop.get("metadata"), dict) else {}
        metadata["postprocessor"] = "mvp11_2_execution_quality_hardening"
        metadata["observed_entry_point"] = observed_entry_point
        agent_sop["metadata"] = metadata

        return agent_sop

    def _compress_detailed_steps(self, steps: list[Any]) -> list[dict[str, Any]]:
        compressed: list[dict[str, Any]] = []
        seen_repeated_keys: set[str] = set()

        for raw_step in steps:
            if not isinstance(raw_step, dict):
                continue
            step = dict(raw_step)
            key = self._detailed_step_key(step)

            if compressed and self._detailed_step_key(compressed[-1]) == key:
                compressed[-1] = self._merge_detailed_step_pair(compressed[-1], step)
                seen_repeated_keys.add(key)
                continue

            # Avoid repeated non-adjacent low-value review/check steps in the same activity.
            if key in seen_repeated_keys and self._is_low_value_repeated_step(step):
                compressed[-1] = self._add_repeat_note_to_detailed_step(compressed[-1])
                continue

            compressed.append(step)

        for index, step in enumerate(compressed, start=1):
            step["step_number"] = index

        return compressed

    def _compress_agent_steps(self, steps: list[Any]) -> list[dict[str, Any]]:
        compressed: list[dict[str, Any]] = []
        seen_keys: set[str] = set()

        for raw_step in steps:
            if not isinstance(raw_step, dict):
                continue
            step = dict(raw_step)
            key = self._agent_step_key(step)

            if compressed and self._agent_step_key(compressed[-1]) == key:
                compressed[-1] = self._merge_agent_step_pair(compressed[-1], step)
                seen_keys.add(key)
                continue

            if key in seen_keys and self._is_low_value_repeated_step(step):
                compressed[-1] = self._add_repeat_note_to_agent_step(compressed[-1])
                continue

            compressed.append(step)
            seen_keys.add(key)

        for index, step in enumerate(compressed, start=1):
            step["step_number"] = index

        return compressed

    def _detailed_step_key(self, step: dict[str, Any]) -> str:
        parts = [
            step.get("instruction", ""),
            step.get("ui_action", ""),
            step.get("action_type", ""),
            step.get("target_ui_element", ""),
            step.get("expected_result", ""),
            step.get("validation", ""),
        ]
        return self._normalize_for_dedupe(" ".join(str(part) for part in parts))

    def _agent_step_key(self, step: dict[str, Any]) -> str:
        parts = [
            step.get("instruction", ""),
            step.get("success_check", ""),
            step.get("action_type", ""),
            step.get("target_ui_element", ""),
        ]
        return self._normalize_for_dedupe(" ".join(str(part) for part in parts))

    def _normalize_for_dedupe(self, text: str) -> str:
        text = self._clean_text(text).lower()
        text = re.sub(r"\bstep\s+\d+\b", "step", text)
        text = re.sub(r"\b\d+\b", "", text)
        text = re.sub(r"[^a-z0-9]+", " ", text)
        text = re.sub(r"\s+", " ", text).strip()
        return text

    def _is_low_value_repeated_step(self, step: dict[str, Any]) -> bool:
        text = " ".join(
            str(step.get(key, ""))
            for key in ["instruction", "ui_action", "expected_result", "validation", "success_check"]
        ).lower()
        repeated_terms = [
            "review", "check", "visible", "displayed", "results", "options", "availability",
            "filter", "scroll", "configure", "preference", "payment options"
        ]
        return any(term in text for term in repeated_terms)

    def _merge_detailed_step_pair(self, left: dict[str, Any], right: dict[str, Any]) -> dict[str, Any]:
        merged = dict(left)
        left_evidence = left.get("evidence") if isinstance(left.get("evidence"), dict) else {}
        right_evidence = right.get("evidence") if isinstance(right.get("evidence"), dict) else {}
        merged_evidence = dict(left_evidence)
        if right_evidence:
            merged_evidence["end_seconds"] = right_evidence.get("end_seconds", merged_evidence.get("end_seconds", 0))
            screen_text = []
            screen_text.extend(left_evidence.get("screen_text", []) or [])
            screen_text.extend(right_evidence.get("screen_text", []) or [])
            merged_evidence["screen_text"] = self._dedupe_keep_order(screen_text)[:8]
            if not merged_evidence.get("frame_path"):
                merged_evidence["frame_path"] = right_evidence.get("frame_path", "")
        merged["evidence"] = merged_evidence
        return self._add_repeat_note_to_detailed_step(merged)

    def _merge_agent_step_pair(self, left: dict[str, Any], right: dict[str, Any]) -> dict[str, Any]:
        merged = dict(left)
        return self._add_repeat_note_to_agent_step(merged)

    def _add_repeat_note_to_detailed_step(self, step: dict[str, Any]) -> dict[str, Any]:
        step = dict(step)
        note = self._clean_text(step.get("fallback_or_note", ""))
        repeat_note = "Repeated similar observations were consolidated. Repeat this action for all visible matching items/options when applicable."
        if repeat_note.lower() not in note.lower():
            step["fallback_or_note"] = f"{note} {repeat_note}".strip()
        return step

    def _add_repeat_note_to_agent_step(self, step: dict[str, Any]) -> dict[str, Any]:
        step = dict(step)
        note = self._clean_text(step.get("notes", ""))
        repeat_note = "Repeat for all visible matching items/options when applicable."
        if repeat_note.lower() not in note.lower():
            step["notes"] = f"{note} {repeat_note}".strip()
        return step


    def _improve_agent_steps_with_observations(
        self,
        steps: list[dict[str, Any]],
        task: dict[str, Any],
        observation_summary: dict[str, Any],
    ) -> list[dict[str, Any]]:
        """Make agent steps more executable using generic observed UI evidence."""

        observed_entry_point = self._clean_text(observation_summary.get("observed_entry_point", ""))
        observed_buttons = observation_summary.get("observed_buttons", []) or []
        observed_input_fields = observation_summary.get("observed_input_fields", []) or []
        observed_result_sections = observation_summary.get("observed_result_sections", []) or []
        observed_form_sections = observation_summary.get("observed_form_sections", []) or []

        improved_steps: list[dict[str, Any]] = []

        for index, raw_step in enumerate(steps, start=1):
            step = dict(raw_step)
            instruction = self._clean_text(step.get("instruction", ""))
            target = self._clean_text(step.get("target_ui_element", ""))
            success_check = self._clean_text(step.get("success_check", ""))
            action_type = self._clean_text(step.get("action_type", ""))

            instruction_lower = instruction.lower()

            if (
                observed_entry_point
                and action_type == "open"
                and (
                    "relevant page" in instruction_lower
                    or "required section" in instruction_lower
                    or "homepage" in instruction_lower
                    or "home page" in instruction_lower
                    or not instruction
                )
            ):
                step["instruction"] = f"Open {observed_entry_point}."
                step["target_ui_element"] = observed_entry_point
                step["success_check"] = (
                    "The observed application entry page is loaded and visible."
                )

            elif action_type in {"type", "enter"} and self._is_vague_ui_text(instruction):
                field_hint = self._first_useful_item(observed_input_fields)
                if field_hint:
                    step["instruction"] = f"Enter the required value in the {field_hint} field."
                    step["target_ui_element"] = field_hint
                    step["success_check"] = f"The {field_hint} field contains the required value."

            elif action_type in {"click", "search", "confirm", "choose", "select"} and self._is_vague_ui_text(target):
                button_hint = self._first_useful_item(observed_buttons)
                if button_hint:
                    step["target_ui_element"] = button_hint

            elif action_type in {"review", "check"} and self._is_vague_ui_text(target):
                result_hint = self._first_useful_item(observed_result_sections)
                if result_hint:
                    step["target_ui_element"] = result_hint

            elif action_type == "select" and self._is_vague_ui_text(instruction):
                form_hint = self._first_useful_item(observed_form_sections)
                if form_hint:
                    step["instruction"] = f"Select the required option in the {form_hint} section."
                    step["target_ui_element"] = form_hint
                    step["success_check"] = f"The selected option is visible in the {form_hint} section."

            if not self._clean_text(step.get("success_check", "")):
                step["success_check"] = success_check or "The expected screen update is visible."

            step["step_number"] = index
            improved_steps.append(step)

        return improved_steps

    def _is_vague_ui_text(self, text: Any) -> bool:
        clean = self._clean_text(text).lower()
        if not clean:
            return True

        vague_phrases = [
            "relevant page",
            "required section",
            "visible criteria",
            "displayed information",
            "visible fields",
            "visible control",
            "visible option",
            "selected option",
            "appropriate option",
            "filter options",
            "additional filter options",
            "additional preferences",
            "additional information",
            "current screen",
            "the page",
            "the field",
            "the button",
            "the option",
        ]

        return any(phrase in clean for phrase in vague_phrases)

    def _first_useful_item(self, items: list[Any]) -> str:
        for item in items:
            clean = self._clean_text(item)
            if clean and len(clean) <= 80:
                return clean
        return ""

    def _build_expected_agent_output(
        self,
        task: dict[str, Any],
        observation_summary: dict[str, Any],
    ) -> list[str]:
        existing = [
            self._clean_text(item)
            for item in task.get("expected_agent_output", []) or []
            if self._clean_text(item)
        ]

        if existing:
            return self._dedupe_keep_order(existing)

        task_text = " ".join(
            [
                self._clean_text(task.get("task_name", "")),
                self._clean_text(task.get("task_purpose", "")),
                " ".join(
                    self._clean_text(step.get("instruction", ""))
                    for step in task.get("steps", []) or []
                    if isinstance(step, dict)
                ),
            ]
        ).lower()

        outputs: list[str] = []

        review_terms = ["review", "check", "search", "result", "results", "list", "availability", "status", "recommend"]
        form_terms = ["enter", "fill", "type", "information", "details", "form"]
        configure_terms = ["configure", "preference", "option", "setting", "select"]

        if any(term in task_text for term in review_terms):
            outputs.extend(
                [
                    "Summarize the visible matching records/options found on the screen.",
                    "Include key visible fields such as name, status, time, class, amount, or other labels shown by the application.",
                    "State any recommendation or selected option with the visible reason.",
                    "Do not invent values that are not visible on screen.",
                ]
            )
        elif any(term in task_text for term in form_terms):
            outputs.extend(
                [
                    "Confirm that all required visible fields were filled.",
                    "List any required fields that remained empty or could not be completed.",
                    "Do not expose sensitive entered values unless required by the user task.",
                ]
            )
        elif any(term in task_text for term in configure_terms):
            outputs.extend(
                [
                    "Summarize the options or preferences selected.",
                    "Confirm that the selected options are visibly applied.",
                    "List any optional settings that were skipped or left unchanged.",
                ]
            )
        else:
            outputs.extend(
                [
                    "Summarize the final visible state reached for this task.",
                    "Mention any visible warning, confirmation, or blocking message.",
                    "Do not claim final business completion unless the completion screen is visible.",
                ]
            )

        return self._dedupe_keep_order(outputs)

    def _clean_completion_criteria(
        self,
        criteria: list[Any],
        source: dict[str, Any],
    ) -> list[str]:
        observed_end_state = self._clean_text(source.get("observed_end_state", ""))
        limitations = " ".join(
            self._clean_text(item)
            for item in source.get("limitations", []) or []
        )
        combined_context = f"{observed_end_state} {limitations}".lower()

        unsafe_final_words = [
            "completed",
            "complete",
            "booked",
            "booking confirmed",
            "paid",
            "payment completed",
            "submitted",
            "approved",
            "deleted",
            "created",
            "saved",
            "finalized",
        ]

        pre_submit_context = any(
            phrase in combined_context
            for phrase in [
                "pre-submit",
                "pre submit",
                "payment-entry",
                "payment entry",
                "review",
                "not observed",
                "final payment",
            ]
        )

        cleaned: list[str] = []
        for item in criteria:
            text = self._clean_text(item)
            if not text:
                continue

            lower = text.lower()
            if pre_submit_context and any(word in lower for word in unsafe_final_words):
                text = (
                    "The workflow reaches the last observed review/pre-submit/payment-entry screen. "
                    "Final submission, payment, approval, save, delete, creation, or completion is not claimed unless visibly confirmed."
                )

            cleaned.append(text)

        if not cleaned and pre_submit_context:
            cleaned.append(
                "The workflow reaches the last observed review/pre-submit/payment-entry screen. Final completion is outside the observed evidence."
            )

        return self._dedupe_keep_order(cleaned)

    def _sop_json_schema(self) -> dict[str, Any]:
        evidence_schema = {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "start_seconds": {"type": "number"},
                "end_seconds": {"type": "number"},
                "speech_summary": {"type": "string"},
                "screen_text": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "frame_path": {"type": "string"},
            },
            "required": [
                "start_seconds",
                "end_seconds",
                "speech_summary",
                "screen_text",
                "frame_path",
            ],
        }

        step_schema = {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "step_number": {"type": "integer"},
                "instruction": {"type": "string"},
                "ui_action": {"type": "string"},
                "action_type": {
                    "type": "string",
                    "enum": [
                        "open",
                        "click",
                        "type",
                        "select",
                        "check",
                        "review",
                        "wait",
                        "scroll",
                        "confirm",
                        "enter",
                        "search",
                        "proceed",
                        "conditional",
                        "other",
                    ],
                },
                "target_ui_element": {"type": "string"},
                "input_value_description": {"type": "string"},
                "expected_result": {"type": "string"},
                "validation": {"type": "string"},
                "condition": {"type": "string"},
                "fallback_or_note": {"type": "string"},
                "evidence": evidence_schema,
            },
            "required": [
                "step_number",
                "instruction",
                "ui_action",
                "action_type",
                "target_ui_element",
                "input_value_description",
                "expected_result",
                "validation",
                "condition",
                "fallback_or_note",
                "evidence",
            ],
        }

        return {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "job_id": {"type": "string"},
                "document_title": {"type": "string"},
                "classification": {"type": "string"},
                "overview": {"type": "string"},
                "scope": {"type": "string"},
                "audience": {"type": "string"},
                "observed_end_state": {"type": "string"},
                "prerequisites": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "general_rules": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "business_process_description": {"type": "string"},
                "activities": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "additionalProperties": False,
                        "properties": {
                            "activity_id": {"type": "string"},
                            "name": {"type": "string"},
                            "description": {"type": "string"},
                            "start_seconds": {"type": "number"},
                            "end_seconds": {"type": "number"},
                            "steps": {
                                "type": "array",
                                "items": step_schema,
                            },
                        },
                        "required": [
                            "activity_id",
                            "name",
                            "description",
                            "start_seconds",
                            "end_seconds",
                            "steps",
                        ],
                    },
                },
                "exceptions_or_notes": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "quality_checks": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "completion_criteria": {
                    "type": "array",
                    "items": {"type": "string"},
                },
                "metadata": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "llm_model": {"type": "string"},
                        "generator": {"type": "string"},
                        "source": {"type": "string"},
                    },
                    "required": [
                        "llm_model",
                        "generator",
                        "source",
                    ],
                },
            },
            "required": [
                "job_id",
                "document_title",
                "classification",
                "overview",
                "scope",
                "audience",
                "observed_end_state",
                "prerequisites",
                "general_rules",
                "business_process_description",
                "activities",
                "exceptions_or_notes",
                "quality_checks",
                "completion_criteria",
                "metadata",
            ],
        }

    def _sop_json_to_markdown(self, sop: dict[str, Any]) -> str:
        lines: list[str] = []

        lines.append(f"# {sop.get('document_title', 'Standard Operating Procedure')}")
        lines.append("")
        lines.append(f"Classification: {sop.get('classification', 'Internal')}")
        lines.append("")

        for title, key in [
            ("Overview", "overview"),
            ("Scope", "scope"),
            ("Audience", "audience"),
            ("Observed End State", "observed_end_state"),
            ("Business Process Description", "business_process_description"),
        ]:
            lines.append(f"## {title}")
            lines.append(str(sop.get(key, "")))
            lines.append("")

        self._append_markdown_list(lines, "Prerequisites", sop.get("prerequisites", []))
        self._append_markdown_list(lines, "General Rules Before Any Action", sop.get("general_rules", []))

        lines.append("## Step-by-Step Procedures")
        lines.append("")

        for activity_index, activity in enumerate(sop.get("activities", []), start=1):
            lines.append(f"### {activity_index}. {activity.get('name', 'Activity')}")
            lines.append("")
            lines.append(activity.get("description", ""))
            lines.append("")

            for step in activity.get("steps", []):
                step_number = step.get("step_number")
                lines.append(f"Step {step_number}: {step.get('instruction', '')}")
                lines.append(f"- Action type: {step.get('action_type', '')}")
                lines.append(f"- Target UI element: {step.get('target_ui_element', '')}")

                input_value = step.get("input_value_description", "")
                if input_value:
                    lines.append(f"- Input value: {input_value}")

                lines.append(f"- Expected result: {step.get('expected_result', '')}")
                lines.append(f"- Validation: {step.get('validation', '')}")

                condition = step.get("condition", "")
                if condition:
                    lines.append(f"- Condition: {condition}")

                note = step.get("fallback_or_note", "")
                if note:
                    lines.append(f"- Note: {note}")

                evidence = step.get("evidence", {})
                start_seconds = evidence.get("start_seconds")
                end_seconds = evidence.get("end_seconds")
                frame_path = evidence.get("frame_path")

                if start_seconds is not None and end_seconds is not None:
                    lines.append(f"- Evidence time: {start_seconds}s - {end_seconds}s")

                if frame_path:
                    lines.append(f"- Frame: `{frame_path}`")

                lines.append("")

        self._append_markdown_list(lines, "Exceptions or Notes", sop.get("exceptions_or_notes", []))
        self._append_markdown_list(lines, "Quality Checks", sop.get("quality_checks", []))
        self._append_markdown_list(lines, "Completion Criteria", sop.get("completion_criteria", []))

        return "\n".join(lines).strip() + "\n"

    def _append_markdown_list(self, lines: list[str], title: str, items: list[Any]) -> None:
        lines.append(f"## {title}")

        if items:
            for item in items:
                lines.append(f"- {item}")
        else:
            lines.append("- None specified.")

        lines.append("")

    def _sop_json_to_docx(self, sop: dict[str, Any], output_path: Path) -> None:
        document = Document()
        self._configure_docx(document)

        classification = sop.get("classification") or "Internal"
        document.add_paragraph(f"Classification: {classification}")

        title = sop.get("document_title") or "Standard Operating Procedure"
        title_paragraph = document.add_paragraph()
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_paragraph.add_run(title)
        title_run.bold = True
        title_run.font.size = Pt(18)

        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_run = subtitle.add_run("Complete Beginner and Agent Execution Guide")
        subtitle_run.italic = True
        subtitle_run.font.size = Pt(11)

        document.add_paragraph("-" * 80)

        self._add_numbered_section(document, 1, "INTRODUCTION", sop.get("overview", ""))
        self._add_numbered_section(document, 2, "SCOPE", sop.get("scope", ""))
        self._add_numbered_section(document, 3, "AUDIENCE", sop.get("audience", ""))
        self._add_numbered_section(document, 4, "OBSERVED END STATE", sop.get("observed_end_state", ""))
        self._add_bullet_section(document, 5, "PREREQUISITES", sop.get("prerequisites", []))
        self._add_bullet_section(document, 6, "GENERAL RULES BEFORE ANY ACTION", sop.get("general_rules", []))
        self._add_numbered_section(
            document,
            7,
            "BUSINESS PROCESS DESCRIPTION",
            sop.get("business_process_description", ""),
        )

        document.add_heading("8. STEP-BY-STEP PROCEDURES", level=1)

        for activity_index, activity in enumerate(sop.get("activities", []), start=1):
            document.add_heading(f"8.{activity_index}. {activity.get('name', 'Activity')}", level=2)

            description = activity.get("description")
            if description:
                document.add_paragraph(str(description))

            for step in activity.get("steps", []):
                step_number = step.get("step_number", "")
                instruction = step.get("instruction", "")

                paragraph = document.add_paragraph()
                run = paragraph.add_run(f"Step {step_number}: {instruction}")
                run.bold = True

                table = document.add_table(rows=0, cols=2)
                table.style = "Table Grid"
                self._add_key_value_row(table, "Action type", step.get("action_type", ""))
                self._add_key_value_row(table, "Target UI element", step.get("target_ui_element", ""))

                input_value = step.get("input_value_description", "")
                if input_value:
                    self._add_key_value_row(table, "Input value", input_value)

                self._add_key_value_row(table, "UI action", step.get("ui_action", ""))
                self._add_key_value_row(table, "Expected result", step.get("expected_result", ""))
                self._add_key_value_row(table, "Validation", step.get("validation", ""))

                condition = step.get("condition", "")
                if condition:
                    self._add_key_value_row(table, "Condition", condition)

                note = step.get("fallback_or_note", "")
                if note:
                    self._add_key_value_row(table, "Fallback / note", note)

                evidence = step.get("evidence", {})
                evidence_text = self._format_evidence(evidence)
                self._add_key_value_row(table, "Evidence", evidence_text)

                document.add_paragraph("")

        self._add_bullet_section(document, 9, "QUALITY CHECKS", sop.get("quality_checks", []))
        self._add_bullet_section(document, 10, "EXCEPTIONS OR NOTES", sop.get("exceptions_or_notes", []))
        self._add_bullet_section(document, 11, "COMPLETION CRITERIA", sop.get("completion_criteria", []))

        document.add_paragraph("-" * 80)
        document.add_paragraph("END OF DOCUMENT")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)

    def _configure_docx(self, document: Document) -> None:
        section = document.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

        styles = document.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(10)

        for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
            style = styles[style_name]
            style.font.name = "Arial"
            style.font.bold = True

        styles["Heading 1"].font.size = Pt(13)
        styles["Heading 2"].font.size = Pt(11)
        styles["Heading 3"].font.size = Pt(10)

    def _add_numbered_section(self, document: Document, number: int, title: str, text: Any) -> None:
        document.add_heading(f"{number}. {title}", level=1)
        document.add_paragraph(str(text or "Not specified."))

    def _add_bullet_section(self, document: Document, number: int, title: str, items: list[Any]) -> None:
        document.add_heading(f"{number}. {title}", level=1)

        if not items:
            document.add_paragraph("None specified.")
            return

        for item in items:
            document.add_paragraph(str(item), style="List Bullet")

    def _add_key_value_row(self, table: Any, key: str, value: Any) -> None:
        row = table.add_row()
        key_cell = row.cells[0]
        value_cell = row.cells[1]

        key_cell.text = str(key)
        value_cell.text = str(value or "")

        if key_cell.paragraphs and key_cell.paragraphs[0].runs:
            key_cell.paragraphs[0].runs[0].bold = True

    def _format_evidence(self, evidence: dict[str, Any]) -> str:
        if not isinstance(evidence, dict):
            return "No evidence reference available."

        parts = []
        start_seconds = evidence.get("start_seconds")
        end_seconds = evidence.get("end_seconds")

        if start_seconds is not None and end_seconds is not None:
            parts.append(f"Time: {start_seconds}s - {end_seconds}s")

        frame_path = evidence.get("frame_path")
        if frame_path:
            parts.append(f"Frame: {frame_path}")

        speech_summary = evidence.get("speech_summary")
        if speech_summary:
            parts.append(f"Speech: {speech_summary}")

        screen_text = evidence.get("screen_text") or []
        if screen_text:
            parts.append("Screen text: " + "; ".join(str(item) for item in screen_text[:6]))

        return " | ".join(parts) if parts else "No evidence reference available."

    def _load_json(self, path: Path) -> dict[str, Any]:
        with path.open("r", encoding="utf-8") as f:
            data = json.load(f)

        if not isinstance(data, dict):
            raise ValueError(f"Expected JSON object in {path}")

        return data

    def _write_json(self, path: Path, data: dict[str, Any]) -> None:
        with path.open("w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

    def _clean_text(self, value: Any) -> str:
        if value is None:
            return ""

        text = str(value)
        return " ".join(text.split())

    def _dedupe_keep_order(self, items: list[Any]) -> list[str]:
        seen = set()
        result = []

        for item in items:
            text = self._clean_text(item)

            if not text:
                continue

            normalized = text.lower()

            if normalized in seen:
                continue

            seen.add(normalized)
            result.append(text)

        return result
